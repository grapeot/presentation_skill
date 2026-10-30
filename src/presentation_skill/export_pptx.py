"""Build an editable PPTX from a JSON slide spec.

The spec keeps text and pictures on separate layers: every word on a slide is a
native PowerPoint text box (so the deck stays editable after a Google Slides
import), while illustrations, diagrams and screenshots are placed as images.
Coordinates are in inches on a 13.333 x 7.5 canvas unless the spec sets
``canvas``.

Inline markup inside text: ``**bold**``, ``~~strike~~`` and ``[[accent]]``
(accent color). Markup does not nest, has no escape, and does not cross a line
break. Every newline starts a new paragraph.
"""

from __future__ import annotations

import json
import re
import shutil
import subprocess
import tempfile
import unicodedata
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

DEFAULT_THEME: dict[str, Any] = {
    "font": "Arial",
    "ea_font": None,  # East Asian face; defaults to "font"
    "text_color": "#1F2A36",
    "muted_color": "#5B6670",
    "accent_color": "#B3261E",
    "background": "#FFFFFF",
    "title_size": 30,
    "body_size": 20,
    "title_box": [0.6, 0.35, 12.1, 0.9],
    "title_bar": None,
    "title_color": None,
    "line_spacing": 1.15,
}

STYLE_PRESETS: dict[str, dict[str, Any]] = {
    "body": {},
    "label": {"bold": True},
    "caption": {"size_delta": -4, "color": "muted"},
    "callout": {"size_delta": -2},
    "small": {"size_delta": -6, "color": "muted"},
}

ALIGN = ("left", "center", "right")
VALIGN = ("top", "middle", "bottom")
FIT = ("contain", "stretch")
DEFAULT_BUBBLE_FILL = "#F3F1EC"
# A line of Arial is about 1.15 em before paragraph spacing is applied.
_FONT_LINE_HEIGHT = 1.15

_MARKUP = re.compile(r"(\*\*.+?\*\*|~~.+?~~|\[\[.+?\]\])")


class SpecError(ValueError):
    """The slide spec is malformed or references a missing file."""


class MissingDependencyError(RuntimeError):
    """python-pptx or an external renderer is not installed."""


class RenderError(RuntimeError):
    """LibreOffice or pdftoppm failed while rendering the deck."""


@dataclass
class BuildResult:
    output: Path
    slides: int
    warnings: list[str] = field(default_factory=list)


def _is_number(value: Any) -> bool:
    return isinstance(value, (int, float)) and not isinstance(value, bool)


def resolve_color(value: Any, theme: dict[str, Any], where: str) -> str:
    """Map a theme color name or #RRGGBB string to #RRGGBB, or raise SpecError."""
    if value in ("text", "muted", "accent"):
        value = theme[f"{value}_color"]
    if not (isinstance(value, str) and re.fullmatch(r"#?[0-9A-Fa-f]{6}", value)):
        raise SpecError(f"{where}: color must be #RRGGBB or one of text, muted, accent; got {value!r}")
    return value if value.startswith("#") else "#" + value


def _rgb(value: str):
    from pptx.dml.color import RGBColor

    return RGBColor.from_string(value.lstrip("#").upper())


def parse_runs(text: str) -> list[tuple[str, dict[str, bool]]]:
    """Split one paragraph into (text, flags) runs using the inline markup."""
    runs: list[tuple[str, dict[str, bool]]] = []
    for piece in _MARKUP.split(text):
        if not piece:
            continue
        if piece.startswith("**") and piece.endswith("**") and len(piece) > 4:
            runs.append((piece[2:-2], {"bold": True}))
        elif piece.startswith("~~") and piece.endswith("~~") and len(piece) > 4:
            runs.append((piece[2:-2], {"strike": True}))
        elif piece.startswith("[[") and piece.endswith("]]") and len(piece) > 4:
            runs.append((piece[2:-2], {"accent": True}))
        else:
            runs.append((piece, {}))
    return runs


def _text_width_em(text: str) -> float:
    """Approximate width in em: full-width (CJK) characters count 1, others 0.5."""
    return sum(1.0 if unicodedata.east_asian_width(ch) in ("W", "F") else 0.5 for ch in text)


def estimate_lines(text: str, width_in: float, size_pt: float, *, bold: bool = False) -> int:
    """Rough line count for Arial-like fonts; used only for overflow warnings."""
    plain = _MARKUP.sub(lambda m: m.group(0)[2:-2], text)
    em_per_line = max(1.0, width_in * 72 / size_pt) / (1.08 if bold else 1.0)
    lines = 0
    for para in plain.split("\n"):
        lines += max(1, int(-(-_text_width_em(para) // em_per_line)))
    return lines


def estimate_height(text: str, width_in: float, size_pt: float, line_spacing: float, *,
                    bold: bool = False) -> float:
    lines = estimate_lines(text, width_in, size_pt, bold=bold)
    return lines * size_pt * _FONT_LINE_HEIGHT * line_spacing / 72


def load_spec(path: Path) -> dict[str, Any]:
    try:
        spec = json.loads(Path(path).read_text(encoding="utf-8"))
    except FileNotFoundError as exc:
        raise SpecError(f"{path}: spec file not found") from exc
    except json.JSONDecodeError as exc:
        raise SpecError(f"{path}: invalid JSON: {exc}") from exc
    if not isinstance(spec, dict):
        raise SpecError(f"{path}: the spec must be a JSON object")
    slides = spec.get("slides")
    if not isinstance(slides, list) or not slides:
        raise SpecError(f"{path}: spec needs a non-empty 'slides' list")
    for i, slide in enumerate(slides, start=1):
        if not isinstance(slide, dict):
            raise SpecError(f"slide {i}: each slide must be a JSON object")
        elements = slide.get("elements", [])
        if not isinstance(elements, list) or not all(isinstance(e, dict) for e in elements):
            raise SpecError(f"slide {i}: 'elements' must be a list of objects")
    canvas = spec.get("canvas", [13.333, 7.5])
    if not (isinstance(canvas, list) and len(canvas) == 2 and all(_is_number(v) and v > 0 for v in canvas)):
        raise SpecError(f"{path}: 'canvas' must be [width, height] in inches")
    if not isinstance(spec.get("theme", {}), dict):
        raise SpecError(f"{path}: 'theme' must be an object")
    return spec


def _box(value: Any, where: str, name: str = "box") -> tuple[float, float, float, float]:
    if not (isinstance(value, list) and len(value) == 4 and all(_is_number(v) for v in value)
            and value[2] > 0 and value[3] > 0):
        raise SpecError(f"{where}: '{name}' must be [x, y, w, h] in inches with positive w and h")
    return tuple(float(v) for v in value)  # type: ignore[return-value]


def _choice(element: dict[str, Any], key: str, allowed: tuple[str, ...], default: str, where: str) -> str:
    value = element.get(key, default)
    if value not in allowed:
        raise SpecError(f"{where}: '{key}' must be one of {', '.join(allowed)}; got {value!r}")
    return value


def _size(value: Any, where: str) -> float:
    if not (_is_number(value) and value > 0):
        raise SpecError(f"{where}: 'size' must be a positive number of points; got {value!r}")
    return float(value)


def _title_bar(value: Any, where: str) -> dict[str, Any] | None:
    if value in (None, False):
        return None
    if not (isinstance(value, dict) and "color" in value):
        raise SpecError(f"{where}: 'title_bar' must be an object with a 'color'")
    height = value.get("height", 1.1)
    if not (_is_number(height) and height > 0):
        raise SpecError(f"{where}: title_bar 'height' must be a positive number")
    return {"color": value["color"], "height": float(height)}


def _fill_text(frame, text: str, *, theme: dict, size: float, color: str, bold: bool,
               italic: bool, align: str, valign: str) -> None:
    from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
    from pptx.oxml.ns import qn
    from pptx.util import Pt

    frame.word_wrap = True
    frame.auto_size = None
    frame.margin_left = frame.margin_right = Pt(4)
    frame.margin_top = frame.margin_bottom = Pt(2)
    frame.vertical_anchor = {"top": MSO_ANCHOR.TOP, "middle": MSO_ANCHOR.MIDDLE,
                             "bottom": MSO_ANCHOR.BOTTOM}[valign]
    ea_font = theme["ea_font"] or theme["font"]
    for i, para_text in enumerate(text.split("\n")):
        para = frame.paragraphs[0] if i == 0 else frame.add_paragraph()
        para.alignment = {"left": PP_ALIGN.LEFT, "center": PP_ALIGN.CENTER,
                          "right": PP_ALIGN.RIGHT}[align]
        para.line_spacing = theme["line_spacing"]
        for run_text, flags in parse_runs(para_text) or [("", {})]:
            run = para.add_run()
            run.text = run_text
            font = run.font
            font.name = theme["font"]
            font.size = Pt(size)
            font.bold = bold or flags.get("bold", False)
            font.italic = italic
            font.color.rgb = _rgb(theme["accent_color"] if flags.get("accent") else color)
            rpr = run._r.get_or_add_rPr()
            if flags.get("strike"):
                rpr.set("strike", "sngStrike")
            for tag, face in (("a:ea", ea_font), ("a:cs", theme["font"])):
                node = rpr.find(qn(tag))
                if node is None:
                    node = rpr.makeelement(qn(tag), {})
                    rpr.append(node)
                node.set("typeface", face)


def _resolve_theme(spec: dict[str, Any]) -> dict[str, Any]:
    theme = {**DEFAULT_THEME, **spec.get("theme", {})}
    for key in ("text_color", "muted_color", "accent_color"):
        theme[key] = resolve_color(theme[key], theme, f"theme.{key}")
    theme["background"] = resolve_color(theme["background"], theme, "theme.background")
    if theme["title_color"] is not None:
        theme["title_color"] = resolve_color(theme["title_color"], theme, "theme.title_color")
    theme["title_bar"] = _title_bar(theme["title_bar"], "theme")
    theme["title_size"] = _size(theme["title_size"], "theme.title_size")
    theme["body_size"] = _size(theme["body_size"], "theme.body_size")
    _box(theme["title_box"], "theme", "title_box")
    if not (_is_number(theme["line_spacing"]) and theme["line_spacing"] > 0):
        raise SpecError("theme.line_spacing must be a positive number")
    for key in ("font", "ea_font"):
        if theme[key] is not None and not (isinstance(theme[key], str) and theme[key]):
            raise SpecError(f"theme.{key} must be a font name")
    return theme


def build_pptx(spec_path: Path, output: Path | None = None) -> BuildResult:
    try:
        from pptx import Presentation
        from pptx.enum.shapes import MSO_SHAPE
        from pptx.util import Inches, Pt
    except ImportError as exc:  # pragma: no cover - exercised only without the extra
        raise MissingDependencyError(
            "python-pptx is required: uv pip install 'presentation-skill[pptx]'"
        ) from exc

    spec_path = Path(spec_path)
    spec = load_spec(spec_path)
    base = spec_path.parent
    theme = _resolve_theme(spec)
    canvas_w, canvas_h = spec.get("canvas", [13.333, 7.5])
    output = Path(output) if output else spec_path.with_suffix(".pptx")
    warnings: list[str] = []

    prs = Presentation()
    prs.slide_width = Inches(canvas_w)
    prs.slide_height = Inches(canvas_h)
    blank = prs.slide_layouts[6]

    for index, slide_spec in enumerate(spec["slides"], start=1):
        where = f"slide {index}"
        slide = prs.slides.add_slide(blank)
        bg = resolve_color(slide_spec.get("background", theme["background"]), theme, f"{where} background")
        slide.background.fill.solid()
        slide.background.fill.fore_color.rgb = _rgb(bg)

        bar = _title_bar(slide_spec["title_bar"], where) if "title_bar" in slide_spec else theme["title_bar"]
        if bar:
            shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, Inches(bar["height"]))
            shape.fill.solid()
            shape.fill.fore_color.rgb = _rgb(resolve_color(bar["color"], theme, f"{where} title_bar"))
            shape.line.fill.background()
            shape.name = "Title bar"

        for element_index, element in enumerate(slide_spec.get("elements", []), start=1):
            ewhere = f"{where} element {element_index}"
            kind = element.get("type")
            if kind == "image":
                _add_image(slide, element, base, ewhere, Inches)
            elif kind in ("text", "bubble"):
                warning = _add_text(slide, element, kind, element_index, theme, ewhere, MSO_SHAPE, Inches, Pt)
                if warning:
                    warnings.append(warning)
            elif kind == "line":
                _add_line(slide, element, element_index, theme, ewhere, Inches, Pt)
            else:
                raise SpecError(f"{ewhere}: unknown element type {kind!r}")

        title = slide_spec.get("title")
        if title:
            if not isinstance(title, str):
                raise SpecError(f"{where}: 'title' must be a string")
            tx, ty, tw, th = _box(slide_spec.get("title_box", theme["title_box"]), where, "title_box")
            box = slide.shapes.add_textbox(Inches(tx), Inches(ty), Inches(tw), Inches(th))
            box.name = "Title"
            title_color = theme["title_color"] or theme["text_color"]
            title_size = _size(slide_spec.get("title_size", theme["title_size"]), f"{where} title_size")
            _fill_text(box.text_frame, title, theme=theme, size=title_size, color=title_color,
                       bold=True, italic=False, align="left", valign="middle")
            needed = estimate_height(title, tw - 0.12, title_size, theme["line_spacing"], bold=True)
            if needed > th + 0.05:
                warnings.append(f"{where}: title may wrap past its box")
        else:
            warnings.append(f"{where}: no title")

        notes = slide_spec.get("notes", "")
        if not isinstance(notes, str):
            raise SpecError(f"{where}: 'notes' must be a string")
        if notes:
            slide.notes_slide.notes_text_frame.text = notes
        else:
            warnings.append(f"{where}: no speaker notes")

    output.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(output))
    return BuildResult(output=output, slides=len(spec["slides"]), warnings=warnings)


def _add_image(slide, element: dict[str, Any], base: Path, where: str, Inches) -> None:
    from PIL import Image, UnidentifiedImageError

    x, y, w, h = _box(element.get("box"), where)
    if not isinstance(element.get("path"), str):
        raise SpecError(f"{where}: image needs a 'path'")
    path = (base / element["path"]).resolve()
    if not path.exists():
        raise SpecError(f"{where}: image not found: {element['path']}")
    fit = _choice(element, "fit", FIT, "contain", where)
    try:
        with Image.open(path) as im:
            ratio = im.width / im.height
    except (UnidentifiedImageError, OSError) as exc:
        raise SpecError(f"{where}: cannot read image {element['path']} (export SVG and other "
                        f"vector formats to PNG first)") from exc
    left, top, pic_w, pic_h = x, y, w, h
    if fit == "contain":
        if w / h > ratio:
            pic_w = h * ratio
            left = x + (w - pic_w) / 2
        else:
            pic_h = w / ratio
            top = y + (h - pic_h) / 2
    pic = slide.shapes.add_picture(str(path), Inches(left), Inches(top), Inches(pic_w), Inches(pic_h))
    pic.name = element.get("name", path.stem)


def _add_text(slide, element: dict[str, Any], kind: str, index: int, theme: dict, where: str,
              MSO_SHAPE, Inches, Pt) -> str | None:
    x, y, w, h = _box(element.get("box"), where)
    style = element.get("style", "body")
    preset = STYLE_PRESETS.get(style)
    if preset is None:
        raise SpecError(f"{where}: unknown style {style!r}; use one of {', '.join(STYLE_PRESETS)}")
    size = _size(element.get("size", theme["body_size"] + preset.get("size_delta", 0)), where)
    color = resolve_color(element.get("color", preset.get("color", "text")), theme, where)
    text = element.get("text", "")
    if not isinstance(text, str):
        raise SpecError(f"{where}: 'text' must be a string")
    align = _choice(element, "align", ALIGN, "left", where)
    valign = _choice(element, "valign", VALIGN, "top", where)
    bold = bool(element.get("bold", preset.get("bold", False)))
    if kind == "bubble":
        shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
        shape.adjustments[0] = 0.12
        shape.fill.solid()
        shape.fill.fore_color.rgb = _rgb(resolve_color(element.get("fill", DEFAULT_BUBBLE_FILL), theme, where))
        if element.get("line"):
            shape.line.color.rgb = _rgb(resolve_color(element["line"], theme, where))
            line_width = element.get("line_width", 1)
            if not (_is_number(line_width) and line_width > 0):
                raise SpecError(f"{where}: 'line_width' must be a positive number")
            shape.line.width = Pt(line_width)
        else:
            shape.line.fill.background()
        shape.shadow.inherit = False
        inner_w = w - 0.25
    else:
        shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        inner_w = w - 0.12
    shape.name = element.get("name", f"{kind} {index}")
    _fill_text(shape.text_frame, text, theme=theme, size=size, color=color, bold=bold,
               italic=bool(element.get("italic", False)), align=align, valign=valign)
    needed = estimate_height(text, max(inner_w, 0.1), size, theme["line_spacing"], bold=bold)
    if text.strip() and needed > h + 0.05:
        return f"{where}: text may overflow ({needed:.2f}in needed, box {h:.2f}in)"
    return None


def _add_line(slide, element: dict[str, Any], index: int, theme: dict, where: str, Inches, Pt) -> None:
    points = element.get("points")
    if not (isinstance(points, list) and len(points) >= 2 and all(
            isinstance(p, list) and len(p) == 2 and all(_is_number(v) for v in p) for p in points)):
        raise SpecError(f"{where}: 'points' needs at least two [x, y] pairs")
    width = element.get("width", 2)
    if not (_is_number(width) and width > 0):
        raise SpecError(f"{where}: 'width' must be a positive number")
    builder = slide.shapes.build_freeform(Inches(points[0][0]), Inches(points[0][1]))
    builder.add_line_segments([(Inches(px), Inches(py)) for px, py in points[1:]], close=False)
    shape = builder.convert_to_shape()
    shape.fill.background()
    shape.line.color.rgb = _rgb(resolve_color(element.get("color", "text"), theme, where))
    shape.line.width = Pt(width)
    shape.name = element.get("name", f"line {index}")


def render_pptx(pptx_path: Path, out_dir: Path, *, dpi: int = 80) -> list[Path]:
    """Render every slide to PNG through LibreOffice and pdftoppm, plus a contact sheet.

    Page files are named by pdftoppm (``slide-1.png`` or zero-padded ``slide-01.png``
    for decks of ten or more slides); the contact sheet is the last returned path.
    """
    soffice = shutil.which("soffice") or shutil.which("libreoffice")
    pdftoppm = shutil.which("pdftoppm")
    if not soffice or not pdftoppm:
        raise MissingDependencyError("rendering needs LibreOffice (soffice) and poppler (pdftoppm) on PATH")
    out_dir = Path(out_dir)
    out_dir.mkdir(parents=True, exist_ok=True)
    with tempfile.TemporaryDirectory() as tmp:
        profile = Path(tmp) / "lo_profile"
        try:
            subprocess.run([soffice, f"-env:UserInstallation={profile.as_uri()}", "--headless",
                            "--convert-to", "pdf", "--outdir", tmp, str(pptx_path)],
                           check=True, capture_output=True, timeout=300)
        except (subprocess.CalledProcessError, subprocess.TimeoutExpired) as exc:
            raise RenderError(f"LibreOffice could not convert {pptx_path}: {exc}") from exc
        pdf = Path(tmp) / (Path(pptx_path).stem + ".pdf")
        if not pdf.exists():
            raise RenderError(f"LibreOffice exited without writing {pdf.name}")
        for old in out_dir.glob("slide-*.png"):
            old.unlink()
        try:
            subprocess.run([pdftoppm, "-r", str(dpi), "-png", str(pdf), str(out_dir / "slide")],
                           check=True, capture_output=True, timeout=300)
        except (subprocess.CalledProcessError, subprocess.TimeoutExpired) as exc:
            raise RenderError(f"pdftoppm failed: {exc}") from exc
    pages = sorted(out_dir.glob("slide-*.png"))
    if not pages:
        raise RenderError("rendering produced no pages")
    return [*pages, _contact_sheet(pages, out_dir / "contact_sheet.png")]


def _contact_sheet(pages: list[Path], target: Path, columns: int = 3) -> Path:
    from PIL import Image, ImageDraw

    thumbs = [Image.open(p).convert("RGB") for p in pages]
    tw = 640
    th = int(thumbs[0].height * tw / thumbs[0].width)
    rows = -(-len(thumbs) // columns)
    gap = 16
    sheet = Image.new("RGB", (columns * (tw + gap) + gap, rows * (th + gap + 22) + gap), "#DDDDDD")
    draw = ImageDraw.Draw(sheet)
    for i, im in enumerate(thumbs):
        r, c = divmod(i, columns)
        x, y = gap + c * (tw + gap), gap + r * (th + gap + 22)
        sheet.paste(im.resize((tw, th)), (x, y + 22))
        draw.text((x, y + 4), f"{i + 1}", fill="#222222")
    sheet.save(target)
    return target
