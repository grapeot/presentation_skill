import json
import shutil
import subprocess
from pathlib import Path

import pytest
from PIL import Image

pptx = pytest.importorskip("pptx")

from presentation_skill.export_pptx import (  # noqa: E402
    SpecError,
    build_pptx,
    estimate_height,
    estimate_lines,
    parse_runs,
    render_pptx,
)

ROOT = Path(__file__).resolve().parents[1]


def _write_spec(tmp_path: Path, slides: list[dict], **extra) -> Path:
    Image.new("RGB", (400, 200), "#336699").save(tmp_path / "plate.png")
    spec = {"slides": slides, **extra}
    path = tmp_path / "deck.json"
    path.write_text(json.dumps(spec), encoding="utf-8")
    return path


def _sample_slides() -> list[dict]:
    return [
        {
            "title": "Results: the effect grows with temperature",
            "elements": [
                {"type": "image", "path": "plate.png", "box": [0.6, 1.4, 6.0, 4.0]},
                {"type": "text", "style": "caption", "box": [0.6, 5.5, 6.0, 0.5],
                 "text": "Figure caption with **bold** and [[accent]]."},
                {"type": "bubble", "box": [7.0, 1.6, 5.5, 1.2], "line": "#999999",
                 "text": "\"check the control\"  ~~old wording~~"},
                {"type": "line", "points": [[7.0, 4.0], [8.5, 3.6], [10.0, 4.2], [12.0, 3.0]],
                 "color": "accent", "width": 3},
            ],
            "notes": "Speaker notes for slide one.",
        },
        {"title": "Second slide", "elements": [], "notes": "More notes."},
    ]


def test_parse_runs_handles_markup():
    runs = parse_runs("plain **bold** mid ~~gone~~ end [[red]]")
    assert ("bold", {"bold": True}) in runs
    assert ("gone", {"strike": True}) in runs
    assert ("red", {"accent": True}) in runs
    assert runs[0] == ("plain ", {})


def test_estimate_lines_grows_with_text():
    assert estimate_lines("short", 6.0, 20) == 1
    assert estimate_lines("word " * 200, 6.0, 20) > 5
    assert estimate_lines("a\nb\nc", 6.0, 20) == 3


def test_build_writes_native_text_images_and_notes(tmp_path: Path):
    spec = _write_spec(tmp_path, _sample_slides(), theme={"font": "Inter"})
    result = build_pptx(spec)
    assert result.output == tmp_path / "deck.pptx"
    assert result.slides == 2

    prs = pptx.Presentation(str(result.output))
    first = prs.slides[0]
    names = [shape.name for shape in first.shapes]
    assert "Title" in names
    texts = " ".join(shape.text_frame.text for shape in first.shapes if shape.has_text_frame)
    assert "check the control" in texts
    assert "Figure caption with bold and accent." in texts
    assert any(shape.shape_type == 13 for shape in first.shapes)  # picture
    assert first.notes_slide.notes_text_frame.text == "Speaker notes for slide one."

    strike_runs = [
        run for shape in first.shapes if shape.has_text_frame
        for para in shape.text_frame.paragraphs for run in para.runs
        if run._r.get_or_add_rPr().get("strike") == "sngStrike"
    ]
    assert [run.text for run in strike_runs] == ["old wording"]
    fonts = {run.font.name for shape in first.shapes if shape.has_text_frame
             for para in shape.text_frame.paragraphs for run in para.runs}
    assert fonts == {"Inter"}


def test_contain_fit_keeps_aspect_ratio(tmp_path: Path):
    spec = _write_spec(tmp_path, [{"title": "t", "notes": "n", "elements": [
        {"type": "image", "path": "plate.png", "box": [1, 1, 4, 4]}]}])
    prs = pptx.Presentation(str(build_pptx(spec).output))
    pic = next(s for s in prs.slides[0].shapes if s.shape_type == 13)
    assert abs(pic.width / pic.height - 2.0) < 0.01


def test_overflow_and_missing_notes_warn(tmp_path: Path):
    spec = _write_spec(tmp_path, [{"title": "t", "elements": [
        {"type": "text", "box": [1, 1, 2, 0.4], "text": "many words " * 40}]}])
    warnings = build_pptx(spec).warnings
    assert any("overflow" in w for w in warnings)
    assert any("no speaker notes" in w for w in warnings)


def test_empty_bubble_does_not_warn(tmp_path: Path):
    spec = _write_spec(tmp_path, [{"title": "t", "notes": "n", "elements": [
        {"type": "bubble", "box": [1, 1, 0.2, 0.2], "text": ""}]}])
    assert build_pptx(spec).warnings == []


@pytest.mark.parametrize(
    "element, message",
    [
        ({"type": "image", "path": "nope.png", "box": [0, 0, 1, 1]}, "image not found"),
        ({"type": "image", "box": [0, 0, 1, 1]}, "needs a 'path'"),
        ({"type": "image", "path": "plate.png", "box": [0, 0, 2, 0]}, "positive w and h"),
        ({"type": "image", "path": "plate.png", "box": [0, 0, 1, 1], "fit": "cover"}, "'fit' must be"),
        ({"type": "image", "path": "fake.svg", "box": [0, 0, 1, 1]}, "cannot read image"),
        ({"type": "text", "box": [0, 0, 1]}, "'box' must be"),
        ({"type": "text", "box": [0, 0, True, 1]}, "'box' must be"),
        ({"type": "shape", "box": [0, 0, 1, 1]}, "unknown element type"),
        ({"type": "text", "style": "fancy", "box": [0, 0, 1, 1]}, "unknown style"),
        ({"type": "text", "box": [0, 0, 1, 1], "color": "purple"}, "color must be"),
        ({"type": "text", "box": [0, 0, 1, 1], "align": "justify"}, "'align' must be"),
        ({"type": "text", "box": [0, 0, 1, 1], "valign": "center"}, "'valign' must be"),
        ({"type": "text", "box": [0, 0, 1, 1], "size": 0}, "'size' must be"),
        ({"type": "bubble", "box": [0, 0, 1, 1], "fill": "#12345"}, "color must be"),
        ({"type": "line", "points": [[0, 0]]}, "at least two"),
        ({"type": "line", "points": [[0, 0], [1]]}, "at least two"),
    ],
)
def test_spec_errors(tmp_path: Path, element: dict, message: str):
    (tmp_path / "fake.svg").write_text("<svg xmlns='http://www.w3.org/2000/svg'/>")
    spec = _write_spec(tmp_path, [{"title": "t", "notes": "n", "elements": [element]}])
    with pytest.raises(SpecError, match=message):
        build_pptx(spec)


@pytest.mark.parametrize(
    "extra, message",
    [
        ({"theme": {"title_bar": {"height": 1}}}, "title_bar"),
        ({"theme": {"title_color": "purple"}}, "color must be"),
        ({"canvas": [13.3]}, "'canvas' must be"),
    ],
)
def test_theme_and_canvas_errors(tmp_path: Path, extra: dict, message: str):
    spec = _write_spec(tmp_path, [{"title": "t", "notes": "n", "elements": []}], **extra)
    with pytest.raises(SpecError, match=message):
        build_pptx(spec)


@pytest.mark.parametrize("content", ['[1, 2]', '{"slides": ["not an object"]}',
                                     '{"slides": [{"elements": ["x"]}]}'])
def test_malformed_structure_rejected(tmp_path: Path, content: str):
    path = tmp_path / "deck.json"
    path.write_text(content)
    with pytest.raises(SpecError):
        build_pptx(path)


def test_missing_spec_file_is_spec_error(tmp_path: Path):
    with pytest.raises(SpecError, match="not found"):
        build_pptx(tmp_path / "missing.json")


def test_theme_color_names_work_everywhere(tmp_path: Path):
    from pptx.dml.color import RGBColor

    spec = _write_spec(tmp_path, [{"title": "t", "notes": "n", "background": "muted", "elements": [
        {"type": "bubble", "box": [1, 1, 3, 1], "text": "x", "fill": "text", "line": "accent"}]}],
        theme={"title_color": "accent", "title_bar": {"color": "muted", "height": 1}})
    prs = pptx.Presentation(str(build_pptx(spec).output))
    slide = prs.slides[0]
    bubble = next(s for s in slide.shapes if s.name == "bubble 1")
    assert bubble.line.color.rgb == RGBColor.from_string("B3261E")
    title = next(s for s in slide.shapes if s.name == "Title")
    assert title.text_frame.paragraphs[0].runs[0].font.color.rgb == RGBColor.from_string("B3261E")
    bar = next(s for s in slide.shapes if s.name == "Title bar")
    assert bar.fill.fore_color.rgb == RGBColor.from_string("5B6670")


def test_theme_fonts_applied_to_latin_and_east_asian(tmp_path: Path):
    from pptx.oxml.ns import qn

    spec = _write_spec(tmp_path, [{"title": "t", "notes": "n", "elements": [
        {"type": "text", "box": [1, 1, 3, 1], "text": "中文 text"}]}],
        theme={"font": "Inter", "ea_font": "Noto Sans SC"})
    prs = pptx.Presentation(str(build_pptx(spec).output))
    run = next(s for s in prs.slides[0].shapes if s.name == "text 1").text_frame.paragraphs[0].runs[0]
    rpr = run._r.get_or_add_rPr()
    assert run.font.name == "Inter"
    assert rpr.find(qn("a:ea")).get("typeface") == "Noto Sans SC"
    assert rpr.find(qn("a:cs")).get("typeface") == "Inter"


def test_contain_centers_tall_image_horizontally(tmp_path: Path):
    Image.new("RGB", (100, 400), "#000000").save(tmp_path / "tall.png")
    spec = _write_spec(tmp_path, [{"title": "t", "notes": "n", "elements": [
        {"type": "image", "path": "tall.png", "box": [1, 1, 4, 4]},
        {"type": "image", "path": "plate.png", "box": [6, 1, 4, 4], "fit": "stretch"}]}])
    prs = pptx.Presentation(str(build_pptx(spec).output))
    tall, stretched = [s for s in prs.slides[0].shapes if s.shape_type == 13]
    emu = 914400
    assert abs(tall.width / emu - 1.0) < 0.01 and abs(tall.left / emu - 2.5) < 0.01
    assert abs(stretched.width / emu - 4) < 0.01 and abs(stretched.height / emu - 4) < 0.01


def test_overflow_estimate_counts_cjk_and_line_height():
    assert estimate_lines("中" * 41, 6.0, 20) == 2
    assert estimate_lines("a" * 41, 6.0, 20) == 1
    assert estimate_height("a\nb\nc\nd", 6.0, 20, 1.15) > 4 * 20 * 1.15 / 72


def test_missing_title_warns(tmp_path: Path):
    spec = _write_spec(tmp_path, [{"notes": "n", "elements": []}])
    assert any("no title" in w for w in build_pptx(spec).warnings)


def test_empty_spec_rejected(tmp_path: Path):
    path = tmp_path / "deck.json"
    path.write_text(json.dumps({"slides": []}))
    with pytest.raises(SpecError):
        build_pptx(path)


def test_cli_build_pptx(tmp_path: Path):
    spec = _write_spec(tmp_path, _sample_slides())
    out = tmp_path / "out" / "deck.pptx"
    result = subprocess.run(
        [str(ROOT / "scripts" / "presentation-skill"), "build-pptx", str(spec), "--output", str(out)],
        cwd=ROOT, text=True, capture_output=True,
    )
    assert result.returncode == 0, result.stdout + result.stderr
    assert out.exists()
    assert "(2 slides)" in result.stdout


def test_cli_build_pptx_bad_spec_exits_2(tmp_path: Path):
    path = tmp_path / "deck.json"
    path.write_text('{"slides": [{"elements": [{"type": "text", "box": [0, 0, 1, 1], "align": "justify"}]}]}')
    result = subprocess.run(
        [str(ROOT / "scripts" / "presentation-skill"), "build-pptx", str(path)],
        cwd=ROOT, text=True, capture_output=True,
    )
    assert result.returncode == 2
    assert "'align' must be" in result.stdout
    assert "Traceback" not in result.stderr


def test_init_pptx_starter_builds(tmp_path: Path):
    out = tmp_path / "deck"
    subprocess.run([str(ROOT / "scripts" / "presentation-skill"), "Thermal transport",
                    "--request", "an editable deck for Google Slides", "--output", str(out)],
                   cwd=ROOT, text=True, capture_output=True, check=True)
    assert "Mode: pptx" in (out / "deck_plan.md").read_text()
    assert build_pptx(out / "slides.json").warnings == []


@pytest.mark.skipif(not (shutil.which("soffice") and shutil.which("pdftoppm")),
                    reason="LibreOffice and pdftoppm are optional")
def test_render_writes_pages_and_contact_sheet(tmp_path: Path):
    spec = _write_spec(tmp_path, _sample_slides())
    pages = render_pptx(build_pptx(spec).output, tmp_path / "render", dpi=30)
    assert [p.name for p in pages] == ["slide-1.png", "slide-2.png", "contact_sheet.png"]


@pytest.mark.skipif(not (shutil.which("soffice") and shutil.which("pdftoppm")),
                    reason="LibreOffice and pdftoppm are optional")
def test_cli_build_pptx_with_render(tmp_path: Path):
    spec = _write_spec(tmp_path, _sample_slides())
    result = subprocess.run(
        [str(ROOT / "scripts" / "presentation-skill"), "build-pptx", str(spec), "--render",
         str(tmp_path / "render"), "--dpi", "30"],
        cwd=ROOT, text=True, capture_output=True,
    )
    assert result.returncode == 0, result.stdout + result.stderr
    assert (tmp_path / "render" / "contact_sheet.png").exists()
